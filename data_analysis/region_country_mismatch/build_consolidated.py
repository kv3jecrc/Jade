import pandas as pd
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

genuine = pd.read_csv("genuine_mismatches.csv", dtype=str, keep_default_na=False)
genuine = genuine.rename(columns={"exp_ship": "Correct Region"})

genuine["Confidence"] = genuine.apply(
    lambda r: "High (Shipping = Billing)" if r["ShippingCountry"] == r["BillingCountry"]
    else "Review (Shipping ≠ Billing)",
    axis=1,
)

genuine["_sort_key"] = genuine["Confidence"].map(
    {"High (Shipping = Billing)": 0, "Review (Shipping ≠ Billing)": 1}
)
genuine = genuine.sort_values(["_sort_key", "Region__c", "Correct Region"]).drop(columns="_sort_key")

select_cols = ["Id", "Name", "ShippingCountry", "BillingCountry", "Region__c", "Correct Region", "Confidence"]
genuine = genuine[select_cols].rename(columns={"Region__c": "Current Region"})
cols = ["Id", "Name", "ShippingCountry", "BillingCountry", "Current Region", "Correct Region", "Confidence"]

sheet_name = "Region Mismatches"
out_path = "consolidated_region_mismatches.xlsx"

with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
    genuine.to_excel(writer, sheet_name=sheet_name, index=False)
    ws = writer.sheets[sheet_name]

    header_font = Font(name="Arial", bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    body_font = Font(name="Arial")
    high_fill = PatternFill(start_color="FCE4E4", end_color="FCE4E4", fill_type="solid")
    review_fill = PatternFill(start_color="FFF6DC", end_color="FFF6DC", fill_type="solid")
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    confidence_col_idx = cols.index("Confidence") + 1
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        confidence_val = row[confidence_col_idx - 1].value
        fill = high_fill if confidence_val and confidence_val.startswith("High") else review_fill
        for cell in row:
            cell.font = body_font
            cell.alignment = Alignment(vertical="center")
            cell.border = border
            cell.fill = fill

    for i, col in enumerate(cols, start=1):
        max_len = max([len(str(col))] + [len(str(v)) for v in genuine[col].astype(str)])
        ws.column_dimensions[get_column_letter(i)].width = min(max_len + 2, 60)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # Summary block below the table
    summary_row = ws.max_row + 3
    ws.cell(row=summary_row, column=1, value="Summary").font = Font(name="Arial", bold=True, size=12)
    high_count = int((genuine["Confidence"].str.startswith("High")).sum())
    review_count = len(genuine) - high_count
    rows = [
        ("Total mismatches listed", len(genuine)),
        ("High confidence (Shipping = Billing country)", high_count),
        ("Needs review (Shipping != Billing country, neither matches Region)", review_count),
    ]
    for offset, (label, value) in enumerate(rows, start=1):
        ws.cell(row=summary_row + offset, column=1, value=label).font = body_font
        ws.cell(row=summary_row + offset, column=2, value=value).font = body_font

print(f"Wrote {out_path}: {len(genuine)} rows ({high_count} high-confidence, {review_count} needing review)")
